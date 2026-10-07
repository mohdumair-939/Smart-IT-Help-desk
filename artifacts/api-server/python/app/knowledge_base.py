"""Document extraction, durable chunks, and local Chroma retrieval."""

import logging
import re
import threading
from io import BytesIO
from pathlib import Path
from uuid import UUID

from docx import Document as DocxDocument
from fastembed import TextEmbedding
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_chroma import Chroma
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import KnowledgeChunk, KnowledgeDocument
from app.object_storage import ObjectStorageError, download_object

logger = logging.getLogger("smart_help_desk.knowledge")
MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
MAX_PDF_PAGES = 250
CHROMA_COLLECTION = "smart_help_desk_knowledge_v1"
CHROMA_PATH = Path(__file__).resolve().parents[2] / ".chroma"
_index_lock = threading.RLock()


class DocumentProcessingError(ValueError):
    pass


class FastEmbedEmbeddings(Embeddings):
    """CPU-friendly local embeddings; source text and vectors stay in the workspace."""

    def __init__(self) -> None:
        self.model = TextEmbedding(model_name="BAAI/bge-small-en-v1.5")

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [vector.tolist() for vector in self.model.embed(texts)]

    def embed_query(self, text: str) -> list[float]:
        return next(self.model.query_embed(text)).tolist()


_embeddings: FastEmbedEmbeddings | None = None
_vector_store: Chroma | None = None


def vector_store() -> Chroma:
    global _embeddings, _vector_store
    with _index_lock:
        if _vector_store is None:
            _embeddings = FastEmbedEmbeddings()
            CHROMA_PATH.mkdir(parents=True, exist_ok=True)
            _vector_store = Chroma(
                collection_name=CHROMA_COLLECTION,
                embedding_function=_embeddings,
                persist_directory=str(CHROMA_PATH),
                collection_metadata={"hnsw:space": "cosine"},
            )
        return _vector_store


def embed_for_similarity(texts: list[str]) -> list[list[float]]:
    with _index_lock:
        vector_store()
        if _embeddings is None:
            raise RuntimeError("The local embedding model was not initialized.")
        return _embeddings.embed_documents(texts)


def embed_query_for_similarity(text: str) -> list[float]:
    with _index_lock:
        vector_store()
        if _embeddings is None:
            raise RuntimeError("The local embedding model was not initialized.")
        return _embeddings.embed_query(text)


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            pdf = PdfReader(BytesIO(content))
            if len(pdf.pages) > MAX_PDF_PAGES:
                raise DocumentProcessingError(f"PDFs may not exceed {MAX_PDF_PAGES} pages.")
            raw_text = "\n\n".join(page.extract_text() or "" for page in pdf.pages)
        elif suffix == ".docx":
            document = DocxDocument(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            )
            raw_text = "\n".join(paragraphs)
        elif suffix in {".txt", ".md"}:
            raw_text = content.decode("utf-8-sig")
        else:
            raise DocumentProcessingError("Upload a PDF, DOCX, TXT, or Markdown document.")
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("The document could not be read. Check that it is a valid, unencrypted file.") from exc

    cleaned = re.sub(r"[ \t]+\n", "\n", raw_text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if len(cleaned) < 20:
        raise DocumentProcessingError("The document contains too little readable text to index.")
    return cleaned


def split_text(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1200,
        chunk_overlap=160,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return [chunk.strip() for chunk in splitter.split_text(text) if chunk.strip()]


def _chunk_document(chunk: KnowledgeChunk) -> Document:
    source = chunk.document
    return Document(
        page_content=chunk.content,
        metadata={
            "document_id": str(source.id),
            "title": source.title,
            "filename": source.filename,
            "chunk_index": chunk.chunk_index,
        },
    )


def ensure_index(db: Session) -> None:
    with _index_lock:
        store = vector_store()
        chunks = db.scalars(
            select(KnowledgeChunk)
            .join(KnowledgeDocument, KnowledgeChunk.document_id == KnowledgeDocument.id)
            .options(selectinload(KnowledgeChunk.document))
            .where(KnowledgeDocument.status == "ready")
        ).all()
        desired: dict[str, Document] = {
            str(chunk.id): _chunk_document(chunk)
            for chunk in chunks
        }
        stored = store.get(include=["metadatas"])
        present_ids = set(stored.get("ids", []))
        stale_ids = present_ids - set(desired)
        if stale_ids:
            store.delete(ids=list(stale_ids))
        missing_ids = list(set(desired) - present_ids)
        if missing_ids:
            store.add_documents(
                documents=[desired[chunk_id] for chunk_id in missing_ids],
                ids=missing_ids,
            )


def add_document_chunks(chunks: list[KnowledgeChunk], db: Session) -> None:
    if not chunks:
        return
    with _index_lock:
        ensure_index(db)
        documents = [_chunk_document(chunk) for chunk in chunks]
        vector_store().add_documents(documents=documents, ids=[str(chunk.id) for chunk in chunks])


def remove_document_chunks(chunk_ids: list[UUID]) -> None:
    if not chunk_ids:
        return
    with _index_lock:
        if _vector_store is not None:
            _vector_store.delete(ids=[str(chunk_id) for chunk_id in chunk_ids])


def retrieve_chunks(db: Session, query: str, limit: int = 5) -> list[tuple[Document, float]]:
    ensure_index(db)
    with _index_lock:
        hits = vector_store().similarity_search_with_relevance_scores(query, k=limit)
    return [(document, float(score)) for document, score in hits if score >= 0.18]


def ingest_document(document_id: UUID, session_factory) -> None:
    """Process a queued document in its own session and leave a visible failure state."""
    db = session_factory()
    pending_chunk_ids: list[UUID] = []
    try:
        document = db.get(KnowledgeDocument, document_id)
        if document is None:
            return
        document.status = "processing"
        document.error_message = None
        db.commit()

        old_chunks = db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document_id)).all()
        old_ids = [chunk.id for chunk in old_chunks]
        remove_document_chunks(old_ids)
        for chunk in old_chunks:
            db.delete(chunk)
        db.flush()

        content = download_object(document.object_path, MAX_DOCUMENT_BYTES)
        text = extract_text(document.filename, content)
        segments = split_text(text)
        if not segments:
            raise DocumentProcessingError("No readable passages were found in the document.")

        chunks = [
            KnowledgeChunk(document_id=document.id, chunk_index=index, content=segment)
            for index, segment in enumerate(segments)
        ]
        pending_chunk_ids = [chunk.id for chunk in chunks]
        db.add_all(chunks)
        document.chunk_count = len(chunks)
        db.flush()
        add_document_chunks(chunks, db)
        document.status = "ready"
        document.error_message = None
        db.commit()
    except (DocumentProcessingError, ObjectStorageError) as exc:
        db.rollback()
        remove_document_chunks(pending_chunk_ids)
        document = db.get(KnowledgeDocument, document_id)
        if document is not None:
            for chunk in db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document_id)).all():
                db.delete(chunk)
            document.status = "error"
            document.chunk_count = 0
            document.error_message = str(exc)[:500]
            db.commit()
        logger.warning("Knowledge document %s could not be indexed: %s", document_id, type(exc).__name__)
    except Exception as exc:
        db.rollback()
        remove_document_chunks(pending_chunk_ids)
        document = db.get(KnowledgeDocument, document_id)
        if document is not None:
            for chunk in db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document_id)).all():
                db.delete(chunk)
            document.status = "error"
            document.chunk_count = 0
            document.error_message = "Indexing failed. Please retry the upload."
            db.commit()
        logger.exception("Unexpected knowledge indexing error (%s)", type(exc).__name__)
    finally:
        db.close()
